import docx
from fpdf import FPDF

# Read the PRD text
with open('BookReviewService_PRD.txt', 'r', encoding='utf-8') as f:
    prd_text = f.read()

# DOCX generation
doc = docx.Document()
for line in prd_text.split('\n'):
    doc.add_paragraph(line)
doc.save('BookReviewService_PRD.docx')

# PDF generation
pdf = FPDF()
pdf.add_page()
pdf.set_auto_page_break(auto=True, margin=15)
pdf.set_font('Arial', '', 12)
for line in prd_text.split('\n'):
    pdf.cell(0, 10, line, ln=True)
pdf.output('BookReviewService_PRD.pdf')
